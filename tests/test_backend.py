from datetime import date
from io import BytesIO
from urllib.parse import urlparse

from backend import main as main_module
from backend.main import app
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import Workbook


def api() -> TestClient:
    return TestClient(app)


def test_student_affairs_dashboard_assets_are_served():
    with api() as client:
        page = client.get("/")
        data = client.get("/static/affairs-data.js")
        app_script = client.get("/static/app.js")

    assert page.status_code == 200
    assert "affairsDashboard" in page.text
    assert "事务助手" in page.text
    assert "/static/affairs-data.js" in page.text
    assert data.status_code == 200
    assert "undergraduate" in data.text
    assert "graduate" in data.text
    assert "官方已核验" in data.text
    assert "往年规律预估" in data.text
    assert "待官方通知" in data.text
    assert "https://cet.neea.edu.cn/" in data.text
    assert 'deadline: "2026-09-18"' not in data.text
    assert 'deadline: "2026-11-15"' not in data.text
    assert app_script.status_code == 200
    assert 'openAssistant("affairs")' in app_script.text
    assert "expectedWindow" in app_script.text


def test_affairs_notices_use_identity_time_and_article_links(monkeypatch):
    async def no_live_news(_profile: str):
        return []

    monkeypatch.setattr(main_module, "fetch_affairs_live_news", no_live_news)
    main_module.AFFAIRS_NEWS_CACHE.clear()
    with api() as client:
        undergraduate = client.get(
            "/api/affairs/notices",
            params={"profile": "undergraduate", "grade": "2024级", "refresh": "true"},
        )
        graduate = client.get(
            "/api/affairs/notices",
            params={"profile": "graduate", "grade": "2025级", "refresh": "true"},
        )

    assert undergraduate.status_code == 200
    assert graduate.status_code == 200
    undergraduate_items = undergraduate.json()["items"]
    graduate_items = graduate.json()["items"]
    assert undergraduate.json()["contentFetchedCount"] == 0
    assert graduate.json()["contentFetchedCount"] == 0
    assert undergraduate_items[0]["publishedDate"] >= undergraduate_items[-1]["publishedDate"]
    assert graduate_items[0]["publishedDate"] >= graduate_items[-1]["publishedDate"]
    allowed_hosts = {"jw.nau.edu.cn", "gs.nau.edu.cn", "xgc.nau.edu.cn", "lib.nau.edu.cn", "cet.neea.edu.cn"}
    assert all(urlparse(item["sourceUrl"]).hostname in allowed_hosts for item in undergraduate_items + graduate_items)
    assert all(urlparse(item["sourceUrl"]).path not in {"", "/"} for item in undergraduate_items + graduate_items)
    assert all("研究生" not in item["matchReason"] for item in undergraduate_items)
    assert all("研究生" in item["matchReason"] for item in graduate_items)


def test_affairs_notice_timeline_validation():
    today = date(2026, 7, 15)
    valid = {
        "publishedDate": "2026-03-18",
        "deadline": "2026-03-30",
        "sourceUrl": "https://jw.nau.edu.cn/2026/0318/c8013a155043/page.htm",
    }
    wrong_url_year = {**valid, "sourceUrl": "https://jw.nau.edu.cn/2025/0318/c8013a155043/page.htm"}
    deadline_before_publication = {**valid, "deadline": "2026-03-10"}
    future_notice = {**valid, "publishedDate": "2026-07-16", "sourceUrl": "https://cet.neea.edu.cn/html1/report/2607/1-1.htm"}

    assert main_module.affairs_news_timeline_is_valid(valid, today) is True
    assert main_module.affairs_news_timeline_is_valid(wrong_url_year, today) is False
    assert main_module.affairs_news_timeline_is_valid(deadline_before_publication, today) is False
    assert main_module.affairs_news_timeline_is_valid(future_notice, today) is False


def test_cet_official_source_can_be_extracted_without_relaxing_host_allowlist():
    source = next(item for item in main_module.AFFAIRS_NEWS_SOURCES if item["name"] == "中国教育考试网·CET")
    html = """
    <li><span>2026-03-06</span>
      <a href="/html1/report/2603/2-1.htm">2026年上半年全国大学英语四、六级考试报名工作启动</a>
      <a href="https://example.com/html1/report/fake.htm">伪造的四六级报名通知不应进入结果</a>
    </li>
    """

    items = main_module.extract_affairs_news(html, source)

    assert len(items) == 1
    assert items[0]["sourceUrl"] == "https://cet.neea.edu.cn/html1/report/2603/2-1.htm"
    assert items[0]["publishedDate"] == "2026-03-06"


def test_affairs_article_body_is_structured_with_deadline_contacts_and_links():
    item = {
        "title": "关于2026年上半年全国大学英语四、六级考试报名的通知",
        "sourceUrl": "https://jw.nau.edu.cn/2026/0318/c8013a155043/page.htm",
    }
    source = next(entry for entry in main_module.AFFAIRS_NEWS_SOURCES if entry["name"] == "教务处")
    html = """
    <html><body><nav>无关导航</nav><article>
      <h1>关于2026年上半年全国大学英语四、六级考试报名的通知</h1>
      <p>发布者：教务处 发布时间：2026-03-18</p>
      <p>本次考试面向全体在校学生（含研究生），请在报名网站完成信息核对。</p>
      <h2>四、正式报名</h2>
      <p>报名时间：2026年3月23日12:00—2026年3月30日17:00，逾期无法补报。</p>
      <p>学生必须在24小时内缴费，如有疑问请联系025-58318571。</p>
      <a href="https://cet-bm.neea.edu.cn/">全国四六级报名系统</a>
      <a href="/_upload/article/files/cet-flow.pdf">附件一：CET报名流程.pdf</a>
    </article><footer>南京审计大学版权所有</footer></body></html>
    """

    details = main_module.extract_affairs_article_details(html, item, source)

    assert details["contentFetched"] is True
    assert details["deadline"] == "2026-03-30"
    assert details["deadlineText"] == "2026年3月30日17:00"
    assert details["audiences"] == ["undergraduate", "graduate"]
    assert "025-58318571" in details["contacts"]
    assert details["serviceLinks"][0]["url"] == "https://cet-bm.neea.edu.cn/"
    assert details["attachments"][0]["url"].endswith("cet-flow.pdf")
    assert any(section["title"] == "正式报名" for section in details["sections"])


def test_affairs_article_parser_filters_template_navigation_links():
    item = {
        "title": "关于2026—2027学年第一学期普通学生教材选购工作的通知",
        "sourceUrl": "https://jw.nau.edu.cn/2026/0710/c8013a157296/page.htm",
    }
    source = next(entry for entry in main_module.AFFAIRS_NEWS_SOURCES if entry["name"] == "教务处")
    html = """
    <html><body><article>
      <h1>关于2026—2027学年第一学期普通学生教材选购工作的通知</h1>
      <p>本科生可在学生个人信息系统选购下学期教材，开放时间为7月12日8:00至7月22日8:00。</p>
      <p>请同学们根据课表安排认真核对教材信息。</p>
      <a href="https://jw.nau.edu.cn/kbcx/list.htm">课表查询</a>
      <a href="https://jw.nau.edu.cn/jwxt/list.htm">教学信息系统</a>
      <a href="https://jw.nau.edu.cn/cxcy/list.htm">创新创业教育平台</a>
      <a href="/_upload/article/files/textbook-notice.pdf">附件：教材选购说明.pdf</a>
    </article></body></html>
    """

    details = main_module.extract_affairs_article_details(html, item, source)

    service_titles = {entry["title"] for entry in details["serviceLinks"]}
    attachment_titles = {entry["title"] for entry in details["attachments"]}
    assert "课表查询" not in service_titles
    assert "教学信息系统" not in service_titles
    assert "创新创业教育平台" not in service_titles
    assert "附件：教材选购说明.pdf" in attachment_titles


def test_chat_creates_session_and_logs_call():
    with api() as client:
        response = client.post(
            "/api/chat",
            json={"mode": "learning", "message": "帮我制定数据结构复习计划"},
        )
        calls = client.get("/api/api-calls").json()

    assert response.status_code == 200
    data = response.json()
    assert data["session_id"]
    assert data["mock"] is True
    assert "学习助手" in data["answer"]
    assert calls
    assert calls[0]["success"] == 1


def test_safety_center_endpoints_create_demo_records():
    with api() as client:
        dashboard = client.get("/api/safety/dashboard")
        sos = client.post(
            "/api/safety/sos",
            json={
                "event_type": "quick help",
                "location": "library first floor",
                "note": "demo only",
                "contact": "student demo",
            },
        )
        report = client.post(
            "/api/safety/reports",
            json={
                "category": "facility risk",
                "description": "light is broken near the dorm gate",
                "location": "Qinyuan dorm gate",
                "urgency": "medium",
                "privacy": "department only",
            },
        )
        trip = client.post(
            "/api/safety/trips",
            json={"destination": "north gate", "contact": "trusted classmate", "duration_minutes": 30},
        )
        lost = client.post(
            "/api/safety/lost-found",
            json={
                "item_type": "丢失",
                "category": "campus card",
                "description": "blue card holder",
                "area": "library area",
            },
        )
        arrive = client.post(f"/api/safety/trips/{trip.json()['id']}/arrive")
        cancelled = client.post(f"/api/safety/sos/{sos.json()['id']}/cancel")

    assert dashboard.status_code == 200
    assert dashboard.json()["campus_status"]["level"]
    assert sos.status_code == 200
    assert sos.json()["status"] == "已发送（演示）"
    assert report.status_code == 200
    assert report.json()["status"] == "待分派"
    assert trip.status_code == 200
    assert trip.json()["status"] == "守护中"
    assert lost.status_code == 200
    assert lost.json()["status"] == "待核验"
    assert arrive.status_code == 200
    assert arrive.json()["status"] == "已安全到达"
    assert cancelled.status_code == 200
    assert cancelled.json()["status"] == "已取消（误触）"


def test_safety_chat_mode_returns_mock_guidance():
    with api() as client:
        response = client.post(
            "/api/chat",
            json={"mode": "safety", "message": "I feel unsafe walking back to dorm alone."},
        )

    assert response.status_code == 200
    assert response.json()["session_id"]
    assert response.json()["mock"] is True
    assert "110" in response.json()["answer"]


def test_second_turn_reuses_history():
    with api() as client:
        first = client.post(
            "/api/chat",
            json={"mode": "learning", "message": "我正在准备智能体竞赛"},
        ).json()
        second = client.post(
            "/api/chat",
            json={
                "mode": "learning",
                "session_id": first["session_id"],
                "message": "那下一步该做什么",
            },
        ).json()
    assert second["session_id"] == first["session_id"]
    assert "延续" in second["answer"]


def test_status_and_documents_available():
    with api() as client:
        status = client.get("/api/status")
    assert status.status_code == 200
    payload = status.json()
    assert payload["document_count"] >= 1
    assert payload["chunk_count"] >= 1


def test_model_config_endpoint_reports_missing_env(monkeypatch):
    monkeypatch.delenv("MODEL_API_KEY", raising=False)
    monkeypatch.delenv("MODEL_BASE_URL", raising=False)
    monkeypatch.delenv("MODEL_NAME", raising=False)

    with api() as client:
        config = client.get("/api/model/config")
        probe = client.post("/api/model/probe", json={})

    assert config.status_code == 200
    assert config.json()["configured"] is False
    assert config.json()["api_key_configured"] is False
    assert probe.status_code == 200
    assert probe.json()["ok"] is False
    assert "MODEL_API_KEY" in probe.json()["error"]


def test_model_config_endpoint_masks_and_normalizes_env(monkeypatch):
    monkeypatch.setenv("MODEL_API_KEY", "sk-test-123456")
    monkeypatch.setenv("MODEL_BASE_URL", "https://api.example.com/v1/chat/completions")
    monkeypatch.setenv("MODEL_NAME", "demo-model")

    with api() as client:
        config = client.get("/api/model/config").json()

    assert config["configured"] is True
    assert config["model_name"] == "demo-model"
    assert config["api_key_tail"] == "3456"
    assert config["base_url"] == "https://api.example.com/v1"
    assert config["chat_completions_url"] == "https://api.example.com/v1/chat/completions"


def test_mock_auth_endpoints():
    with api() as client:
        status = client.get("/api/auth/status")
        login = client.post(
            "/api/auth/mock-login",
            json={"username": "20260001", "password": "demo"},
        )
        logout = client.post("/api/auth/logout")

    assert status.status_code == 200
    assert status.json()["portal_url"] == "https://my.nau.edu.cn/index.html#/"
    assert login.status_code == 200
    assert login.json()["authenticated"] is True
    assert logout.status_code == 200
    assert logout.json()["authenticated"] is False


def test_stream_returns_expected_sse_events():
    with api() as client:
        with client.stream(
            "POST",
            "/api/chat/stream",
            json={"mode": "teaching", "message": "生成一份课堂活动设计"},
        ) as response:
            body = response.read().decode("utf-8")
    assert response.status_code == 200
    assert "event: meta" in body
    assert "event: token" in body
    assert "event: references" in body
    assert "event: done" in body


def test_deep_research_and_web_search_modes():
    with api() as client:
        response = client.post(
            "/api/chat",
            json={
                "mode": "audit",
                "message": "分析审计证据可靠性",
                "deep_research": True,
                "web_search": True,
            },
        )

    assert response.status_code == 200
    assert "深度研究" in response.json()["answer"]
    assert "未执行真实网络检索" in response.json()["answer"]


def test_upload_txt_and_docx_generate_chunks():
    with api() as client:
        txt_response = client.post(
            "/api/upload",
            files={"file": ("course.txt", "数据结构课程包含线性表、树、图、排序等知识点。".encode("utf-8"), "text/plain")},
        )

        doc = Document()
        doc.add_paragraph("教学设计应包含教学目标、课堂活动、试题和 Rubric。")
        doc_bytes = BytesIO()
        doc.save(doc_bytes)
        doc_bytes.seek(0)
        docx_response = client.post(
            "/api/upload",
            files={
                "file": (
                    "teaching.docx",
                    doc_bytes.getvalue(),
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
        )

    assert txt_response.status_code == 200
    assert txt_response.json()["chunk_count"] >= 1
    assert docx_response.status_code == 200
    assert docx_response.json()["chunk_count"] >= 1


def test_report_generation():
    with api() as client:
        chat = client.post(
            "/api/chat",
            json={"mode": "affairs", "message": "奖学金怎么申请"},
        ).json()
        report = client.post("/api/reports", json={"session_id": chat["session_id"]})
    assert report.status_code == 200
    assert report.json()["download_url"].startswith("/api/reports/")


def test_audit_mode_and_artifacts():
    with api() as client:
        chat = client.post(
            "/api/chat",
            json={"mode": "audit", "message": "解释审计风险模型"},
        )
        mindmap = client.post(
            "/api/artifacts/mindmap",
            json={"mode": "audit", "prompt": "审计风险模型"},
        )
        quiz = client.post(
            "/api/artifacts/quiz",
            json={"mode": "audit", "prompt": "审计证据", "question_count": 4},
        )
        workpaper = client.post(
            "/api/artifacts/audit-workpaper",
            json={"mode": "audit", "prompt": "政府采购合规性案例"},
        )

    assert chat.status_code == 200
    assert "审计学习助手" in chat.json()["answer"]
    assert mindmap.status_code == 200
    assert "mindmap" in mindmap.json()["content"]
    assert quiz.status_code == 200
    assert "测验题" in quiz.json()["content"]
    assert workpaper.status_code == 200
    assert "审计工作底稿" in workpaper.json()["content"]


def test_upload_xlsx_generate_chunks():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "报销明细"
    sheet.append(["日期", "人员", "金额", "事项"])
    sheet.append(["2026-01-05", "张三", 1888, "差旅费"])
    sheet.append(["2026-01-06", "张三", 1888, "差旅费"])
    payload = BytesIO()
    workbook.save(payload)
    payload.seek(0)

    with api() as client:
        response = client.post(
            "/api/upload",
            files={
                "file": (
                    "audit-data.xlsx",
                    payload.getvalue(),
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )

    assert response.status_code == 200
    assert response.json()["chunk_count"] >= 1


def test_translation_text_and_file_export():
    with api() as client:
        text = client.post(
            "/api/translate/text",
            json={"source_text": "audit evidence and internal control", "target_language": "中文"},
        )
        exported = client.post(
            "/api/translate/file",
            data={"target_language": "中文"},
            files={"file": ("audit-note.txt", b"audit evidence and working paper", "text/plain")},
        )
        summary = client.post(
            "/api/translate/summary",
            json={"source_text": "audit evidence and internal control", "target_language": "中英双语"},
        )
        download = client.get(exported.json()["download_url"])

    assert text.status_code == 200
    assert "审计证据" in text.json()["translation"]
    assert exported.status_code == 200
    assert exported.json()["download_url"].startswith("/api/exports/")
    assert summary.status_code == 200
    assert "English Key Points" in summary.json()["summary"]
    assert download.status_code == 200


def test_document_history_supports_preview_download_and_images():
    with api() as client:
        uploaded = client.post(
            "/api/upload",
            files={"file": ("notes.txt", "审计证据学习笔记".encode("utf-8"), "text/plain")},
        ).json()
        image = client.post(
            "/api/upload",
            files={"file": ("course-photo.png", b"\x89PNG\r\n\x1a\n", "image/png")},
        )
        details = client.get(f"/api/documents/{uploaded['document_id']}")
        download = client.get(uploaded["download_url"])

    assert image.status_code == 200
    assert image.json()["indexed"] is False
    assert details.status_code == 200
    assert "审计证据学习笔记" in details.json()["content"]
    assert download.status_code == 200


def test_artifact_history_and_archive_to_knowledge_base():
    with api() as client:
        generated = client.post(
            "/api/artifacts/outline",
            json={"mode": "audit", "prompt": "审计风险模型"},
        ).json()
        history = client.get("/api/artifacts")
        archived = client.post(
            f"/api/artifacts/{generated['artifact_id']}/archive",
            json={},
        )

    assert history.status_code == 200
    assert any(item["id"] == generated["artifact_id"] for item in history.json())
    assert archived.status_code == 200
    assert archived.json()["scope"] == "generated"


def test_import_timetable_from_xlsx():
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["课程名称", "教师", "星期", "节次", "教室", "周次"])
    sheet.append(["审计学基础", "王老师", "星期一", "1-2节", "敏达楼101", "1-16周"])
    payload = BytesIO()
    workbook.save(payload)
    payload.seek(0)

    with api() as client:
        imported = client.post(
            "/api/timetable/import",
            files={
                "file": (
                    "课程表.xlsx",
                    payload.getvalue(),
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )
        timetable = client.get("/api/timetable")

    assert imported.status_code == 200
    assert imported.json()["imported_count"] == 1
    assert timetable.status_code == 200
    assert any(item["course_name"] == "审计学基础" for item in timetable.json())
